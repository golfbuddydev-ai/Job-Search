"""
Step 4 -- tailored CVs for the roles you pick.

How to use: in the Matches tab, type Y in the "Apply?" column next to any
role you want a CV for. Then run "4 - Tailor CVs" from the Actions tab
(it also runs once a day by itself, and costs nothing if no new rows are
marked Y).

For each role marked Y it:
  1. Reads your full Master Profile (CV + everything from the interviews)
  2. Has Claude rewrite your REAL experience for that role -- reordering
     and rewording, never inventing experience, numbers or titles
  3. Builds a Word document (.docx)
  4. Attaches it to the run in GitHub Actions, and puts a link to that run
     in the "Tailored CV Link" column so you can download it

Always read the CV before you send it.
"""

import os
import json
import datetime

from anthropic import Anthropic
from docx import Document
from docx.shared import Pt, Inches

import config
import sheets_helper
from claude_helper import extract_text

PROFILE_HEADER = ["Section", "Content", "Last Updated", "Source"]
APPLY_COL = "Apply?"
LINK_COL = "Tailored CV Link"
OUTPUT_DIR = "tailored_cvs"

LABEL_INSTRUCTION = """
FORMATTING: every bullet point must start with a short skill/competency label (2-4 words) \
naming the capability that bullet demonstrates, then the detail -- e.g. label "Client Growth", \
text "Grew the top 10 accounts by 30% in a year...". Use the "label" field for it."""

PLAIN_INSTRUCTION = """
FORMATTING: plain bullet points. Leave the "label" field as an empty string."""


def generate_tailored_content(client, profile, job):
    system_prompt = f"""Here is a candidate's full master profile (their real career history --
use ONLY what's here, never invent experience, numbers, or skills not present):

{profile}

Write a tailored CV for this candidate, aimed at this specific role:
Title: {job.get('Title', '')}
Company: {job.get('Company', '')}
Why it's a fit: {job.get('Why It Fits', '')}

Emphasize and reorder the candidate's REAL experience to speak directly to this role --
reword bullet points to use language relevant to the role, lead with the most relevant
experience, and de-emphasize (but don't necessarily remove) less relevant history. Do not
fabricate metrics, titles, dates, or responsibilities that aren't grounded in the profile above.
{LABEL_INSTRUCTION if config.CV_BULLET_LABELS else PLAIN_INSTRUCTION}

Respond with ONLY a JSON object, no other text, in this exact shape:
{{
  "name": "...",
  "contact_line": "...",
  "summary": "2-3 sentences tailored to this role",
  "experience": [
    {{"role_company_dates": "...", "bullets": [
      {{"label": "...", "text": "..."}}
    ]}}
  ],
  "skills": "comma-separated, most relevant to this role first",
  "education": "..."
}}"""

    response = client.messages.create(
        model=config.CLAUDE_MODEL,
        max_tokens=4000,
        system=system_prompt,
        messages=[{"role": "user", "content": "Generate the tailored CV content."}],
    )
    raw = extract_text(response).replace("```json", "").replace("```", "").strip()
    return json.loads(raw)


def _heading(doc, text):
    doc.add_paragraph()
    doc.add_paragraph().add_run(text).bold = True


def build_docx(content, output_path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.8)
        section.top_margin = section.bottom_margin = Inches(0.7)

    name_run = doc.add_paragraph().add_run(content.get("name", ""))
    name_run.bold = True
    name_run.font.size = Pt(20)
    doc.add_paragraph(content.get("contact_line", ""))

    _heading(doc, "SUMMARY")
    doc.add_paragraph(content.get("summary", ""))

    _heading(doc, "EXPERIENCE")
    for role in content.get("experience", []):
        doc.add_paragraph().add_run(role.get("role_company_dates", "")).bold = True
        for bullet in role.get("bullets", []):
            p = doc.add_paragraph(style="List Bullet")
            if isinstance(bullet, dict):
                label = (bullet.get("label") or "").strip()
                if label:
                    p.add_run(f"{label}: ").bold = True
                p.add_run(bullet.get("text", ""))
            else:
                p.add_run(str(bullet))

    _heading(doc, "SKILLS")
    doc.add_paragraph(content.get("skills", ""))

    _heading(doc, "EDUCATION")
    doc.add_paragraph(content.get("education", ""))

    doc.save(output_path)


def get_run_url():
    server = os.environ.get("GITHUB_SERVER_URL", "https://github.com")
    repo = os.environ.get("GITHUB_REPOSITORY", "")
    run_id = os.environ.get("GITHUB_RUN_ID", "")
    if repo and run_id:
        return f"{server}/{repo}/actions/runs/{run_id}"
    return f"(run on your computer -- see the {OUTPUT_DIR}/ folder)"


def main():
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise SystemExit("ANTHROPIC_API_KEY is not set. See SETUP.md, step 5.")

    gc = sheets_helper.get_client()
    matches_tab = sheets_helper.open_tab(
        gc, config.MATCHES_TAB, ["Date Found", "Fit Score", "Title", "Company", "Link", APPLY_COL, LINK_COL]
    )
    link_col = matches_tab.row_values(1).index(LINK_COL) + 1

    to_do = [
        (i, row) for i, row in enumerate(matches_tab.get_all_records(), start=2)
        if str(row.get(APPLY_COL, "")).strip().upper() in ("Y", "YES")
        and not str(row.get(LINK_COL, "")).strip()
    ]
    if not to_do:
        print("No new roles marked Y in the 'Apply?' column -- nothing to do (and nothing spent).")
        return

    profile_rows = sheets_helper.open_tab(gc, config.PROFILE_TAB, PROFILE_HEADER).get_all_records()
    if not profile_rows:
        raise SystemExit("Your Master Profile is empty -- run '1 - Build profile from CV' first.")
    profile = sheets_helper.profile_text(profile_rows)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    client = Anthropic(api_key=api_key)
    done = 0

    for i, row in to_do:
        print(f"Writing a CV for: {row.get('Title', '')} at {row.get('Company', '')}...")
        try:
            content = generate_tailored_content(client, profile, row)
        except Exception as e:
            print(f"  [warn] couldn't generate this one ({e}) -- it'll be retried next run.")
            continue

        safe_company = "".join(c if c.isalnum() else "_" for c in str(row.get("Company", "role")))[:40]
        filename = f"CV_{safe_company}_{datetime.date.today().isoformat()}_row{i}.docx"
        build_docx(content, os.path.join(OUTPUT_DIR, filename))
        matches_tab.update_cell(i, link_col, get_run_url())
        print(f"  Saved {filename}")
        done += 1

    print(f"Done. Wrote {done} tailored CV(s). Download them from this run's page, "
          "under 'Artifacts' at the bottom.")


if __name__ == "__main__":
    main()
